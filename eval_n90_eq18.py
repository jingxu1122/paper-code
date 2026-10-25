# -*- coding: utf-8 -*-
"""Evaluate LSDI and comparison indices on the 90 field spectra.

Input
  D:\\prosail\\2019_72_2026_18_n90_eq18.xlsx
  Spectral values are read as stored.

Protocols
  (A) PROSAIL train, field test. LSDI bands come from the
      simulation spectra only.
  (B) Field stratified 70/30, 50 repeats; report mean and standard deviation.
      LSDI bands are the same as in (A).
"""
from __future__ import annotations

import json
import pickle
import shutil
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook
from scipy.stats import rankdata
from sklearn.metrics import confusion_matrix

import litter_index_2  # noqa: F401  Times New Roman
from litter_index_2 import (
    ALL_COVERAGES, XLSX_MEAN_PATH, SIM_EXPORT_PKL, OUT_DIR,
    SEARCH_RANGE, MIN_SEP, step1_read_measured_mean, step3_accuracy_filter,
    calc_base_vi, BASE_VI,
)
from plot_all_sg_figures import sg_mat, ORDER_VI, fit_predict, scores
from rf_7030_field import split_7030, fit_predict as fit_predict_7030, class_rates, RF_TREES

FIELD_XLSX = Path(r'D:\prosail\2019_72_2026_18_n90_eq18.xlsx')
OUT = Path(__file__).resolve().parent
TAG = 'n90_eq18'
N_7030 = 50
NAMES_VI = ['LSDI'] + [v for v in ORDER_VI if v in BASE_VI]

def _spearman_abs(mat, y):
    rx = np.apply_along_axis(rankdata, 0, mat)
    ry = rankdata(y).astype(float)
    rx = rx - rx.mean(0)
    ry = ry - ry.mean()
    den = np.sqrt((rx ** 2).sum(0) * (ry ** 2).sum())
    den = np.where(den < 1e-12, np.nan, den)
    return np.abs((rx * ry[:, None]).sum(0) / den)


def lsdi_from_bands(spec_mat, ia, ib, ic):
    ba = spec_mat[:, ia].astype(float)
    bb = spec_mat[:, ib].astype(float)
    bc = spec_mat[:, ic].astype(float)
    return (ba - bb) / (ba + bb + bc + 1e-10)



def load_field_xlsx(path, wl):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    names = [str(c) for c in rows[0][1:]]
    data = np.array([[np.nan if v is None else v for v in r] for r in rows[1:]], float)
    if not np.allclose(data[:, 0], wl, atol=0.5):
        raise RuntimeError('wavelength grid does not match reference')
    X = data[:, 1:].T.astype(float)
    y_cov = []
    for n in names:
        if 'LU2' in n:
            y_cov.append(100)
            continue
        cov = None
        for p in n.split('_'):
            if p.endswith('%'):
                cov = int(p[:-1])
                break
        if cov is None:
            raise RuntimeError(f'cannot parse coverage from {n}')
        y_cov.append(cov)
    return names, X, np.array(y_cov, int)


def lsdi_bands_from_simulation(X_sim, y_cov_sim, wl, band_idx, chunk=4000):
    B = X_sim[:, band_idx].astype(float)
    wlc = wl[band_idx]
    nb = B.shape[1]
    sep = np.abs(wlc[:, None] - wlc[None, :]) >= MIN_SEP
    iu = np.triu_indices(nb, k=1)
    keep = sep[iu]
    ii, jj = iu[0][keep], iu[1][keep]
    best_score, best_pair = -1.0, (0, 1)
    for start in range(0, len(ii), chunk):
        end = min(start + chunk, len(ii))
        ba, bb = B[:, ii[start:end]], B[:, jj[start:end]]
        vi = (ba - bb) / (np.abs(ba) + np.abs(bb) + 1e-10)
        sc = _spearman_abs(vi, y_cov_sim)
        sc = np.where(np.all(np.isfinite(vi), axis=0), sc, -1.0)
        k = int(np.nanargmax(sc))
        if sc[k] > best_score:
            best_score = float(sc[k])
            best_pair = (int(ii[start + k]), int(jj[start + k]))
    ia = int(band_idx[best_pair[0]])
    ib = int(band_idx[best_pair[1]])
    ba, bb = X_sim[:, ia].astype(float), X_sim[:, ib].astype(float)
    ic_list, cols = [], []
    for ic in band_idx:
        if ic in (ia, ib):
            continue
        if abs(wl[ic] - wl[ia]) < MIN_SEP or abs(wl[ic] - wl[ib]) < MIN_SEP:
            continue
        bc = X_sim[:, ic].astype(float)
        denom = ba + bb + bc
        lsdi = (ba - bb) / (denom + 1e-10)
        if np.any(np.abs(denom) < 1e-10) or not np.all(np.isfinite(lsdi)):
            continue
        ic_list.append(int(ic))
        cols.append(lsdi)
    scores3 = _spearman_abs(np.column_stack(cols), y_cov_sim)
    ic = ic_list[int(np.nanargmax(scores3))]
    return ia, ib, ic


def eval_7030_locked(X, y_cov, y, wl, names_vi, n_rep, ia, ib, ic):
    X_sg = sg_mat(X)
    fixed = calc_base_vi(X_sg, wl)
    feats = {'LSDI': lsdi_from_bands(X_sg, ia, ib, ic), **fixed}
    store = {n: [] for n in names_vi}
    cms = {n: [] for n in names_vi}
    for r in range(n_rep):
        tr, te = split_7030(y, seed=r)
        for name in names_vi:
            yhat = fit_predict_7030(feats[name][tr], y[tr], feats[name][te])
            oa, _, _, cm = class_rates(y[te], yhat)
            store[name].append(oa)
            cms[name].append(cm)
    means = {n: float(np.mean(store[n])) for n in names_vi}
    stds = {n: float(np.std(store[n], ddof=1)) if n_rep > 1 else 0.0 for n in names_vi}
    mean_cm = {n: np.mean(np.stack(cms[n], axis=0), axis=0) for n in names_vi}
    return means, stds, mean_cm


def plot_oa_bars(names, accs, title, path, stds=None):
    order = np.argsort(accs)[::-1]
    names_s = [names[i] for i in order]
    accs_s = [accs[i] for i in order]
    stds_s = [stds[i] for i in order] if stds is not None else None
    clrs = ['#C0392B' if n == 'LSDI' else '#BDC3C7' for n in names_s]
    fig, ax = plt.subplots(figsize=(4.8, 2.95) if stds_s else (4.6, 2.85))
    x = np.arange(len(names_s))
    if stds_s is None:
        bars = ax.bar(x, accs_s, 0.62, color=clrs, edgecolor='white',
                      linewidth=0.4, zorder=3)
        for bar, v in zip(bars, accs_s):
            ax.text(bar.get_x() + bar.get_width() / 2, v + 0.012, f'{v:.0%}',
                    ha='center', va='bottom', fontsize=6.5, color='#2C3E50')
        ax.set_ylim(0.0, min(1.0, max(accs_s) + 0.16))
        ax.set_ylabel('Classification accuracy', fontsize=8)
    else:
        bars = ax.bar(x, accs_s, 0.62, yerr=stds_s, color=clrs, edgecolor='white',
                      linewidth=0.4, zorder=3, capsize=2,
                      error_kw=dict(lw=0.7, ecolor='0.35'))
        for bar, v, s in zip(bars, accs_s, stds_s):
            ax.text(bar.get_x() + bar.get_width() / 2, min(1.08, v + s + 0.018),
                    f'{v:.0%}', ha='center', va='bottom', fontsize=6.5, color='#2C3E50')
        ax.set_ylim(0.0, 1.18)
        ax.set_ylabel('Test accuracy', fontsize=8)
    ax.set_xticks(x)
    ax.set_xticklabels(names_s, fontsize=7.5, rotation=35, ha='right')
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f'{v:.0%}'))
    ax.yaxis.grid(True, lw=0.4, color='0.88', zorder=0)
    ax.set_axisbelow(True)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.legend(handles=[
        Patch(facecolor='#C0392B', edgecolor='none', label='LSDI'),
        Patch(facecolor='#BDC3C7', edgecolor='none', label='Comparison indices'),
    ], fontsize=6.5, loc='upper right', handlelength=1.0, borderpad=0.4)
    ax.text(0.0, 1.02, title, transform=ax.transAxes, fontsize=6,
            va='bottom', ha='left', color='0.35')
    fig.tight_layout(pad=0.55)
    fig.savefig(path, dpi=600, bbox_inches='tight')
    plt.close()
    print(f'  saved {path}')


def plot_confusion(y_true, pred, acc, names, path, subtitle):
    cmap_cm = LinearSegmentedColormap.from_list('cm_blue', ['#FDFEFE', '#1A5276'], N=256)
    n_col = 5
    n_row = int(np.ceil(len(names) / n_col))
    fig, axes = plt.subplots(n_row, n_col, figsize=(8.8, 4.15),
                             gridspec_kw={'wspace': 0.42, 'hspace': 0.62})
    axes = np.atleast_1d(axes).ravel()
    labels = ['Low', 'High']
    for k, name in enumerate(names):
        ax = axes[k]
        cm = confusion_matrix(y_true, pred[name], labels=[0, 1])
        vmax = cm.max() if cm.max() else 1
        ax.imshow(cm, cmap=cmap_cm, aspect='equal', vmin=0, vmax=vmax)
        thresh = vmax / 2.0
        for i in range(2):
            row_sum = cm[i].sum()
            for j in range(2):
                val = cm[i, j]
                pct = val / row_sum * 100 if row_sum else 0.0
                ax.text(j, i, f'{val}\n({pct:.0f}%)', ha='center', va='center',
                        fontsize=7, color='white' if val > thresh else '#2C3E50',
                        linespacing=1.25)
        ax.set_xticks([0, 1])
        ax.set_yticks([0, 1])
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_xlabel('Predicted', fontsize=8)
        if k % n_col == 0:
            ax.set_yticklabels(labels, fontsize=8)
            ax.set_ylabel('True label', fontsize=8)
        else:
            ax.set_yticklabels([''] * 2)
        ax.set_title(f'{name}  OA={acc[name]:.0%}', fontsize=9, pad=4)
        ax.tick_params(length=0)
        for spine in ax.spines.values():
            spine.set_visible(False)
    for k in range(len(names), len(axes)):
        fig.delaxes(axes[k])
    fig.suptitle(subtitle, fontsize=9, color='0.25', y=1.02)
    fig.tight_layout(pad=0.6)
    fig.savefig(path, dpi=600, bbox_inches='tight')
    plt.close()
    print(f'  saved {path}')


def plot_mean_confusion(mean_cm, acc, names, path, subtitle):
    """Mean cell counts over repeated 70/30 splits. Percentages are row shares of the mean matrix."""
    cmap_cm = LinearSegmentedColormap.from_list('cm_blue', ['#FDFEFE', '#1A5276'], N=256)
    n_col = 5
    n_row = int(np.ceil(len(names) / n_col))
    fig, axes = plt.subplots(n_row, n_col, figsize=(8.8, 4.15),
                             gridspec_kw={'wspace': 0.42, 'hspace': 0.62})
    axes = np.atleast_1d(axes).ravel()
    labels = ['Low', 'High']
    for k, name in enumerate(names):
        ax = axes[k]
        cm = np.asarray(mean_cm[name], float)
        vmax = cm.max() if cm.max() else 1
        ax.imshow(cm, cmap=cmap_cm, aspect='equal', vmin=0, vmax=vmax)
        thresh = vmax / 2.0
        for i in range(2):
            row_sum = cm[i].sum()
            for j in range(2):
                val = cm[i, j]
                pct = val / row_sum * 100 if row_sum else 0.0
                ax.text(j, i, f'{val:.1f}\n({pct:.0f}%)', ha='center', va='center',
                        fontsize=7, color='white' if val > thresh else '#2C3E50',
                        linespacing=1.25)
        ax.set_xticks([0, 1])
        ax.set_yticks([0, 1])
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_xlabel('Predicted', fontsize=8)
        if k % n_col == 0:
            ax.set_yticklabels(labels, fontsize=8)
            ax.set_ylabel('True label', fontsize=8)
        else:
            ax.set_yticklabels([''] * 2)
        ax.set_title(f'{name}  OA={acc[name]:.0%}', fontsize=9, pad=4)
        ax.tick_params(length=0)
        for spine in ax.spines.values():
            spine.set_visible(False)
    for k in range(len(names), len(axes)):
        fig.delaxes(axes[k])
    fig.suptitle(subtitle, fontsize=9, color='0.25', y=1.02)
    fig.tight_layout(pad=0.6)
    fig.savefig(path, dpi=600, bbox_inches='tight')
    plt.close()
    print(f'  saved {path}')


def main():
    wl, measured_mean = step1_read_measured_mean(XLSX_MEAN_PATH)
    names, X, y_cov = load_field_xlsx(FIELD_XLSX, wl)
    y = (y_cov >= 60).astype(int)
    print(f'Field spectra: {FIELD_XLSX.name}')
    print(f'  n={len(y)}  high(≥60%)={int(y.sum())}  low(<60%)={int((y == 0).sum())}')
    print(f'  coverage {dict(sorted(Counter(y_cov.tolist()).items()))}')

    tmp = Path(OUT_DIR) / f'_step3_tmp_{TAG}'
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    with open(SIM_EXPORT_PKL, 'rb') as f:
        sim_data = pickle.load(f)['sim_data']
    sim_data, _, _, _ = step3_accuracy_filter(wl, measured_mean, sim_data, out_dir=tmp)
    shutil.rmtree(tmp, ignore_errors=True)

    X_sim = np.concatenate([sim_data[c]['all'] for c in ALL_COVERAGES])
    y_cov_sim = np.concatenate([np.full(len(sim_data[c]['all']), c) for c in ALL_COVERAGES])
    band_idx = np.where((wl >= SEARCH_RANGE[0]) & (wl <= SEARCH_RANGE[1]))[0]
    ia, ib, ic = lsdi_bands_from_simulation(X_sim, y_cov_sim, wl, band_idx)
    bands = (int(round(wl[ia])), int(round(wl[ib])), int(round(wl[ic])))
    print(f'LSDI bands (simulation): {bands}')

    # ---- (A) PROSAIL → field ----
    X_tr = sg_mat(X_sim)
    y_tr = (y_cov_sim >= 60).astype(int)
    X_te = sg_mat(X)
    vi_tr = calc_base_vi(X_tr, wl)
    vi_te = calc_base_vi(X_te, wl)
    vi_tr['LSDI'] = lsdi_from_bands(X_tr, ia, ib, ic)
    vi_te['LSDI'] = lsdi_from_bands(X_te, ia, ib, ic)

    print(f'\n(A) PROSAIL → field  |  RF trees=200')
    pred_ext, acc_ext, rows_ext = {}, {}, []
    sim_accs = []
    for name in NAMES_VI:
        yhat = fit_predict(vi_tr[name], y_tr, vi_te[name])
        sc = scores(y, yhat)
        pred_ext[name] = yhat
        acc_ext[name] = sc['oa']
        sim_accs.append(sc['oa'])
        rows_ext.append(dict(
            Index=name, Protocol='PROSAIL→field SG',
            OA=round(sc['oa'], 4), Low=round(sc['low'], 4), High=round(sc['high'], 4),
        ))
        print(f'  {name:<8} OA={sc["oa"]:.3f}  Low={sc["low"]:.3f}  High={sc["high"]:.3f}')
    pd.DataFrame(rows_ext).to_csv(
        OUT / f'external_test_{TAG}_SG.csv', index=False, encoding='utf-8-sig')
    plot_oa_bars(
        NAMES_VI, sim_accs,
        f'$n$ = {len(y)} field spectra   |   Binary (≥60% vs. <60%)',
        OUT / f'Fig_external_test_{TAG}_SG.png',
    )
    plot_confusion(
        y, pred_ext, acc_ext, NAMES_VI,
        OUT / f'Fig_external_test_{TAG}_SG_confusion.png',
        f'$n$ = {len(y)} field spectra   |   Binary (≥60% vs. <60%)',
    )

    # ---- (B) field 70/30 ----
    print(f'\n(B) field 70/30  |  {N_7030} repeats, mean±std  |  RF trees={RF_TREES}')
    f_means, f_stds, f_cm = eval_7030_locked(X, y_cov, y, wl, NAMES_VI, N_7030, ia, ib, ic)
    rows_f = []
    for name in NAMES_VI:
        rows_f.append(dict(
            Index=name, Protocol='field 70/30 SG',
            OA_mean=round(f_means[name], 4), OA_std=round(f_stds[name], 4),
        ))
        print(f'  {name:<8} {f_means[name]:.3f}±{f_stds[name]:.3f}')
    pd.DataFrame(rows_f).sort_values('OA_mean', ascending=False).to_csv(
        OUT / f'rf_7030_{TAG}.csv', index=False, encoding='utf-8-sig')
    plot_oa_bars(
        NAMES_VI, [f_means[n] for n in NAMES_VI],
        f'Train 63 / test 27   |   n = {len(y)}',
        OUT / f'Fig_rf_7030_{TAG}.png',
        stds=[f_stds[n] for n in NAMES_VI],
    )

    plot_mean_confusion(
        f_cm, f_means, NAMES_VI,
        OUT / f'Fig_rf_7030_{TAG}_confusion_mean.png',
        f'Train 63 / test 27   |   Binary (≥60% vs. <60%)',
    )

    meta = dict(
        field_xlsx=str(FIELD_XLSX),
        n_samples=int(len(y)),
        coverage=dict(sorted(Counter(y_cov.tolist()).items())),
        lsdi_bands_nm=list(bands),
        external_OA={n: round(acc_ext[n], 4) for n in NAMES_VI},
        field7030_OA_mean={n: round(f_means[n], 4) for n in NAMES_VI},
        field7030_OA_std={n: round(f_stds[n], 4) for n in NAMES_VI},
        n_7030_repeats=N_7030,
    )
    with open(OUT / f'{TAG}_meta.json', 'w', encoding='utf-8') as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    print(f'  saved {OUT / f"{TAG}_meta.json"}')
    print('done')


if __name__ == '__main__':
    main()
